#!/usr/bin/env python3
"""Planilha de orçamento dos equipamentos de produção da EJUP.

Preços de lojas brasileiras pesquisados em 09/10/2026 (itens novos). Para cada
item há o preço médio e o menor preço visto em loja séria; quantidade e preços
são entradas (azul), subtotais e totais são fórmulas. Depois de rodar,
recalcular com o recalc.py da skill xlsx.
"""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = pathlib.Path(__file__).resolve().parents[1] / "Orcamento_Equipamentos_EJUP.xlsx"
DATA = "09/10/2026"

C, A, L, S, E = "Câmeras e lentes", "Áudio", "Iluminação", "Suporte e energia", "Edição"
# (categoria, item, qtd, preço médio, menor preço visto, onde comprar, link, observação)
ITENS = [
    (C, "Sony FX30 (corpo)", 2, 10800, 9774, "Bella Photo (oficial, garantia Sony Brasil), KaBuM, Amazon",
     "https://www.bellaphoto.com.br/camera-sony-fx30-digital-cinema",
     "Na Bella: R$ 9.809 no Pix. No marketplace, garantia de 3 meses indica importação paralela."),
    (C, "Sony A7 IV (corpo)", 1, 14600, 12579, "Horizon Play, KaBuM, Magalu",
     "https://www.horizonplay.com.br/produto/camera-digital-sony-a7-iv-ilce-7m4-33mp-3-0-lente-fe-28-70mm-oss-corpo-150843",
     "Está sumindo das lojas. A sucessora, A7 V, sai por R$ 16.199 no Pix na Bella Photo."),
    (C, "Tamron 28-75mm f/2.8 Di III VXD G2 (Sony E)", 2, 6100, 5400, "Detona Shop, Bella Photo, Magalu",
     "https://www.bellaphoto.com.br/lente-tamron-28-75mm-f28-di-iii-vxd-g2-para-cameras-sony-e",
     "Detona informa garantia Tamron Brasil de 1 ano. Oferta da Amazon com garantia dos EUA é importada."),
    (C, "Tamron 16-30mm f/2.8 Di III VXD G2 (Sony E), a lente aberta", 1, 6200, 5849, "Bella Photo",
     "https://www.bellaphoto.com.br/lente-tamron-16-30mm-f28-di-iii-vxd-g2-para-cameras-sony-e",
     "Substitui a 17-28 f/2.8, que saiu de linha. Custa cerca de R$ 450 a mais e é a sucessora direta."),
    (C, "Bateria Sony NP-FZ100 original", 3, 650, 589, "Shop AV / Loja dos Marios, Bella Photo",
     "https://www.shop-av.com/bateria-sony-np-fz100-cameras-a7iii-a6600",
     "Serve nas três câmeras. Com as 3 que vêm nas caixas, ficam 6."),
    (C, "Bateria falsa NP-FZ100 com fonte de tomada (bivolt)", 3, 205, 172, "WorldView, Amazon, Magalu",
     "https://www.worldview.com.br/adaptador-ac-pw20-com-acoplador-dc-np-fz100-para-bateria-sony-fz100-bivolt/p",
     "Para gravação longa sem trocar bateria. Testar dentro do prazo de troca: há relato de falha após atualização de firmware."),
    (C, "Carregador duplo NP-FZ100 (SmallRig 4087)", 1, 230, 123, "Shop AV, Magalu",
     "https://www.shop-av.com/produto/smallrig-4087-carregador-de-bateria-sony-np-fz100.html",
     "Na Shop AV o prazo de entrega é de 20 dias úteis."),
    (C, "Cartão SD Lexar 128 GB V90 UHS-II (2000x Gold)", 3, 1050, 750, "Shop AV, Amazon, Magalu",
     "https://www.shop-av.com/produto/cartao-de-memoria-sdxc-lexar-2000x-128-gb-v90-uhs-ii-300mbs.html",
     "V90 grava qualquer formato da FX30, inclusive o All-Intra. Preços muito baixos no marketplace pedem checagem do vendedor."),
    (C, "Mattebox SmallRig Mini Matte Box Lite (3575)", 1, 600, 600, "Mercado Livre (compra internacional)",
     "https://lista.mercadolivre.com.br/matt-box",
     "Encaixa nas lentes de 67 mm com anel adaptador: confirmar se o anúncio traz o anel. Alternativa robusta: Tilta Mini MB-T15, R$ 1.100 a 1.183."),

    (A, "Hollyland Lark M3, kit 2 transmissores + 2 receptores", 1, 999, 999, "Amazon (único lugar com estoque)",
     "https://www.amazon.com.br/hollyland-lark/s?k=hollyland+lark",
     "O kit de 1 receptor só tem receptor USB-C e não liga na câmera. Por isso o de 2 receptores."),
    (A, "Zoom PodTrak P4 (ou P4next)", 1, 1500, 1199, "Mercado Livre, Amazon, Magalu",
     "https://www.mercadolivre.com.br/gravador-de-podcast-multifaixa-portatil-zoom-podtrak-p4-cor-cinza/p/MLB27949896",
     "O P4 está saindo de linha. O P4next (R$ 1.443 na Amazon) grava em 24 bits, mas usa microSD e fonte de 5V/2,4A."),
    (A, "Fonte USB 5V com cabo USB-C (para o gravador)", 1, 55, 37, "Mercado Livre, Amazon",
     "https://lista.mercadolivre.com.br/fonte-5v-2a-usb-c",
     "O P4 vem sem cabo. Preferir marca conhecida para não gerar ruído. Para o P4next, 5V/2,4A."),
    (A, "Cartão SD Lexar 64 GB V30 (para o gravador)", 1, 150, 107, "Mercado Livre, KaBuM, Amazon",
     "https://www.mercadolivre.com.br/carto-de-memoria-lexar-sd-xc-64gb-150mbs-uhs-i-v30-pro/p/MLB37831026",
     "O P4 usa SD de tamanho normal. Se for o P4next, trocar por microSD do mesmo tamanho."),
    (A, "Microfone dinâmico Fifine K688 (XLR/USB)", 4, 470, 390, "Pichau, Terabyte, KaBuM",
     "https://www.kabum.com.br/produto/435116/microfone-dinamico-fifine-k688-usb-xlr-preto",
     "No lugar do BM-800: capta menos o som da sala. Alternativas: Samson Q2U (~R$ 780, já com cabo) ou Shure SM58 (~R$ 1.100)."),
    (A, "Suporte de mesa tipo tripé para microfone", 4, 55, 47, "Magalu, Amazon",
     "https://www.bondfaro.com.br/pedestal-para-microfone?page=2",
     "Como já é feito nos podcasts. O K688 costuma vir sem suporte: conferir no anúncio. Rosca 5/8\" com adaptador para 3/8\"."),
    (A, "Cabo XLR macho-fêmea 3 m (Santo Angelo)", 4, 60, 46, "Amazon, Magalu, Mercado Livre",
     "https://www.amazon.com.br/MICROFONE-CONECTOR-MACHO-FEMEA-METROS/dp/B076N55W36",
     "Se o microfone escolhido for o Samson Q2U, não precisa: ele já vem com cabo."),
    (A, "Fone JBL Tune 500 com fio", 1, 135, 116, "Mercado Livre, Amazon, KaBuM",
     "https://www.zoom.com.br/fone-de-ouvido-e-headset/headphone-com-microfone-jbl-tune-500",
     "Com fio, para ouvir o gravador sem atraso. Versão sem fio (Tune 500BT) não serve para monitorar."),

    (L, "Amaran 200d S (no lugar do Aputure LS 120D II)", 1, 3100, 2662, "Mercado Livre, Amazon",
     "https://www.mercadolivre.com.br/luz-de-video-led-diurna-aputure-amaran-200d-s-200w-5600k/p/MLB2001538692",
     "O 120D II saiu de linha. O 200d S tem encaixe Bowens e mais potência. Preferir anúncio com nota fiscal."),
    (L, "Octabox Triopo K2-120 (120 cm) com colmeia, encaixe Bowens", 1, 640, 497, "Amazon, Mercado Livre, Shopee",
     "https://www.amazon.com.br/Softbox-Octabox-Triopo-K2-120-Encaixe/dp/B0DY5SGMB4",
     "O kit da Amazon traz 2 difusores, colmeia e bolsa. Há anúncios sem colmeia: conferir."),
    (L, "Tripé com girafa Greika CS-330 (C-stand)", 1, 1000, 999, "Amazon, Magalu, Mercado Livre",
     "https://www.amazon.com.br/Trip%C3%A9-Girafa-Ilumina%C3%A7%C3%A3o-C-Stand-330cm/dp/B0FH7P78NW",
     "Carga de 10 kg com a haste. É o que segura a luz e o octabox sobre a mesa."),
    (L, "Saco de areia para tripé", 2, 65, 31, "Mercado Livre, Amazon",
     "https://lista.mercadolivre.com.br/saco-de-areia-tripe",
     "Em geral vem vazio. Segurança da luz pendurada sobre a mesa."),
    (L, "Bastão LED Yongnuo YN360 III, kit com fonte 12V", 1, 680, 646, "Amazon",
     "https://www.amazon.com.br/Bast%C3%A3o-Yongnuo-YN360-Video-Bivolt/dp/B07TDN2HRS",
     "O kit com fonte sai mais barato que bastão e fonte separados. Usa 1 bateria NP-F."),
    (L, "Tripé de luz simples Greika (cerca de 2 m)", 1, 90, 49, "Mercado Livre, Amazon",
     "https://www.amazon.com.br/Trip%C3%A9-Ilumina%C3%A7%C3%A3o-Greika-00-PK-SB03/dp/B07XLTYY6F",
     "Para o bastão."),
    (L, "Bateria NP-F970 (Greika ou Tudoprafoto)", 2, 220, 142, "Amazon, Magalu",
     "https://www.magazineluiza.com.br/bateria-digital-video-sony-np-f970-greika/p/hgc922h285/cf/bpfi/",
     "Para o bastão. Opção premium: SmallRig com USB-C, cerca de R$ 900."),
    (L, "Carregador para bateria NP-F", 1, 120, 76, "Magalu, Mercado Livre",
     "https://www.magazineluiza.com.br/carregador-de-bateria-sony-e-similares-np-f970-tudoprafoto/p/6041920/cf/bpfi/",
     "Carregador duplo SmallRig: cerca de R$ 208."),

    (S, "Tripé de vídeo Ulanzi Ombra Video II (T154), cabeça fluida", 3, 600, 575, "Amazon, loja oficial Ulanzi no Mercado Livre",
     "https://www.amazon.com.br/ULANZI-libera%C3%A7%C3%A3o-port%C3%A1til-telefone-grava%C3%A7%C3%A3o/dp/B0F5WSY485",
     "Carga de 3 kg, 157 cm. Para mais folga: Ulanzi MT-56 (6 kg), cerca de R$ 800."),
    (S, "Mesa dobrável portátil tipo maleta (alumínio)", 1, 160, 107, "Mercado Livre, Amazon, Leroy Merlin",
     "https://www.amazon.com.br/Mesa-Dobr%C3%A1vel-Alum%C3%ADnio-695838-Palisad/dp/B0777QNWGL",
     "Tamanhos de 60x45 a 120x60 cm."),
    (S, "Filtro de linha iClamper Energia 8 tomadas", 2, 105, 94, "KaBuM, Terabyte, Pichau",
     "https://www.kabum.com.br/produto/616087/filtro-de-linha-dps-iclamper-energia-8-8-tomadas-bivolt-preto-26775",
     "Com proteção contra surto."),
    (S, "Extensão elétrica 5 m, 3 tomadas", 2, 57, 51, "Amazon, Leroy Merlin, Magalu",
     "https://www.amazon.com.br/Tramontina-Extens%C3%A3o-M%C3%BAltipla-Tomadas-250/dp/B076X58SD9",
     "Para várias cargas, preferir cabo de 1,0 mm² ou mais."),
    (E, "MacBook Pro 14\" M1 Pro, 16 GB, 512 GB, usado em ótimo estado", 1, 7100, 5880,
     "Mac Switch e GGImports (seminovo com garantia); OLX e Mercado Livre (particular)",
     "https://www.macswitch.com.br/products/macbook-pro-m1-pro-14",
     "Seminovo com garantia de loja: R$ 6.900 a 7.500. Particular: R$ 5.000 a 9.500, sem garantia. Conferir saúde da bateria, ciclos, conta Apple desvinculada e carregador original."),
    (E, "HD externo Seagate Expansion 4 TB (USB 3.0)", 1, 1190, 1115, "KaBuM, Amazon",
     "https://www.kabum.com.br/produto/1005358/hd-externo-portatil-seagate-expansion-4tb-2-5-usb-3-0-preto-stgx4000400",
     "Para arquivo e backup. Para editar direto, o SSD é melhor."),
    (E, "SSD externo SanDisk Portable 2 TB (800 MB/s, USB-C)", 1, 1900, 1859, "KaBuM",
     "https://www.kabum.com.br/produto/513298/ssd-externo-sandisk-2tb-800mbs-usb-portatil-grafite",
     "No lugar do SanDisk de 560 MB/s: aquele é SSD interno SATA, precisa de case USB-C e sai mais caro (R$ 2.000 a 2.400). Este é externo, mais rápido e já liga no MacBook."),
]

LOJAS = [
    ("Bella Photo", "Câmeras, lentes e baterias Sony",
     "A melhor verificada: estoque real, 10% de desconto no Pix e 1 ano de garantia. Na FX30, garantia Sony Brasil."),
    ("Detona Shop", "Câmeras Sony e lentes Tamron",
     "Diz ser revenda oficial Sony, com 12 meses de garantia, e garantia Tamron Brasil nas lentes. O site estava instável no dia."),
    ("Shop AV / Loja dos Marios", "Baterias, carregadores e cartões",
     "A mais barata em acessórios, com nota fiscal. Prazo de entrega de 9 a 20 dias úteis."),
    ("KaBuM, Magalu e Amazon", "Quase tudo",
     "São marketplaces: olhar sempre o \"vendido por\". Garantia \"do fornecedor\" de 3 meses indica importação paralela."),
    ("Pichau, Terabyte e KaBuM", "Microfones, braços e fones",
     "Revenda nacional com nota fiscal. Fifine com estoque no Brasil."),
    ("Mundomax e Playtech", "Áudio profissional (Shure, Audio-Technica)",
     "Lojas de música com garantia oficial (24 meses no Shure da Mundomax)."),
    ("Amazon", "Lark M3, octabox Triopo, Greika, Ulanzi, kit YN360",
     "Único lugar com o Lark M3 hoje. Conferir o vendedor e a garantia."),
    ("Mercado Livre", "Luz amaran, mattebox, acessórios, loja oficial Ulanzi",
     "Muitos anúncios são compra internacional: preço menor, mas sem garantia nacional e com prazo maior. Preferir anúncio com nota fiscal."),
    ("Leroy Merlin e KaBuM", "Mesa, filtro de linha e extensão", "Preços estáveis, produto nacional."),
    ("Mac Switch e GGImports", "MacBook seminovo",
     "Lojas de seminovos Apple com garantia (3 meses na Mac Switch) e troca em 7 dias. Particular na OLX sai mais barato, mas sem garantia: só comprar vendo a máquina ligada."),
    ("Importar dos EUA (B&H, Adorama)", "Não compensa",
     "Fora do Remessa Conforme: 60% de imposto mais ICMS. Uma FX30 de US$ 1.750 chegaria a cerca de R$ 18 mil, contra R$ 9,8 mil na Bella Photo."),
    ("Trazer na mala", "Pode compensar em lente",
     "Cota de US$ 1.000; o excedente paga 50%. A Tamron 28-75 cabe na cota e economiza de R$ 1.000 a 1.500. Perde a garantia nacional."),
]

AZUL = Font(name="Arial", size=10, color="0000FF")
PRETO = Font(name="Arial", size=10)
NEG = Font(name="Arial", size=10, bold=True)
TIT = Font(name="Arial", size=14, bold=True, color="031C37")
CAB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FUNDO_CAB = PatternFill("solid", fgColor="031C37")
FUNDO_CAT = PatternFill("solid", fgColor="E9EDF2")
FUNDO_TOT = PatternFill("solid", fgColor="F6EFD9")
BORDA = Border(bottom=Side(style="thin", color="C8CDD3"))
MOEDA = 'R$ #,##0;-R$ #,##0;"-"'
TOPO = Alignment(vertical="top", wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = "Orçamento"
ws["A1"] = "Orçamento de equipamentos · produção EJUP"
ws["A1"].font = TIT
ws["A2"] = (f"Preços de produtos novos em lojas brasileiras, pesquisados em {DATA}. Preços mudam todo dia: "
            "conferir vendedor, garantia e estoque antes de fechar.")
ws["A2"].font = Font(name="Arial", size=9, italic=True, color="5F6062")
ws["A3"] = "Azul: quantidade e preços, que podem ser editados. Preto: fórmulas."
ws["A3"].font = Font(name="Arial", size=9, color="5F6062")

cab = ["Categoria", "Item", "Qtd", "Preço médio (R$)", "Subtotal médio (R$)", "Menor preço visto (R$)",
       "Subtotal menor (R$)", "Onde comprar", "Link", "Observação"]
for c, t in enumerate(cab, 1):
    cel = ws.cell(row=5, column=c, value=t)
    cel.font, cel.fill = CAB, FUNDO_CAB
    cel.alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[5].height = 30

lin = 6
cats = []
for cat, item, qtd, medio, menor, onde, link, obs in ITENS:
    if cat not in cats:
        cats.append(cat)
    vals = {1: cat, 2: item, 8: onde, 10: obs}
    for c, v in vals.items():
        ws.cell(row=lin, column=c, value=v).font = PRETO
    ws.cell(row=lin, column=3, value=qtd).font = AZUL
    for c, v in ((4, medio), (6, menor)):
        cel = ws.cell(row=lin, column=c, value=v)
        cel.font, cel.number_format = AZUL, MOEDA
    for c, f in ((5, f"=C{lin}*D{lin}"), (7, f"=C{lin}*F{lin}")):
        cel = ws.cell(row=lin, column=c, value=f)
        cel.font, cel.number_format = PRETO, MOEDA
    lk = ws.cell(row=lin, column=9, value=link)
    lk.hyperlink = link
    lk.font = Font(name="Arial", size=9, color="304557", underline="single")
    for c in range(1, 11):
        ws.cell(row=lin, column=c).border = BORDA
        ws.cell(row=lin, column=c).alignment = TOPO
    lin += 1
ult = lin - 1

lin += 1
ws.cell(row=lin, column=2, value="Total por categoria").font = NEG
lin += 1
for cat in cats:
    ws.cell(row=lin, column=2, value=cat).font = PRETO
    for col, src in ((5, "E"), (7, "G")):
        t = ws.cell(row=lin, column=col, value=f'=SUMIF($A$6:$A${ult},B{lin},${src}$6:${src}${ult})')
        t.font, t.number_format = PRETO, MOEDA
    for c in range(1, 11):
        ws.cell(row=lin, column=c).fill = FUNDO_CAT
    lin += 1
ws.cell(row=lin, column=2, value="TOTAL GERAL").font = NEG
for col, src in ((5, "E"), (7, "G")):
    t = ws.cell(row=lin, column=col, value=f"=SUM({src}6:{src}{ult})")
    t.font, t.number_format = NEG, MOEDA
for c in range(1, 11):
    ws.cell(row=lin, column=c).fill = FUNDO_TOT
lin += 2
ws.cell(row=lin, column=2, value=("Subtotal médio: média das lojas sérias pesquisadas. Subtotal menor: o menor preço visto, "
                                  "quase sempre no Pix ou em marketplace; serve de piso para negociação.")).font = Font(
    name="Arial", size=9, italic=True, color="5F6062")

for col, w in zip("ABCDEFGHIJ", (15, 44, 6, 14, 15, 14, 15, 30, 40, 58)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "C6"

wl = wb.create_sheet("Onde comprar")
wl["A1"] = "Onde comprar"
wl["A1"].font = TIT
for c, t in enumerate(["Loja", "Para quê", "Observação"], 1):
    cel = wl.cell(row=3, column=c, value=t)
    cel.font, cel.fill = CAB, FUNDO_CAB
for i, linha in enumerate(LOJAS, 4):
    for c, v in enumerate(linha, 1):
        cel = wl.cell(row=i, column=c, value=v)
        cel.font = NEG if c == 1 else PRETO
        cel.alignment, cel.border = TOPO, BORDA
for col, w in zip("ABC", (30, 42, 90)):
    wl.column_dimensions[col].width = w

# Excel, Google Sheets e Numbers recalculam tudo ao abrir
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print(OUT)


def gravar_valores(caminho):
    """Grava o resultado de cada fórmula como valor em cache, para visualizadores que não recalculam.
    As fórmulas continuam no arquivo, e o Excel recalcula ao abrir (fullCalcOnLoad)."""
    import re, zipfile, shutil, tempfile
    vals = {}
    for i, (cat, item, q, m, n, *_ ) in enumerate(ITENS):
        r = 6 + i
        vals[f"E{r}"], vals[f"G{r}"] = q * m, q * n
    ult = 5 + len(ITENS)
    r = ult + 3
    for cat in cats:
        vals[f"E{r}"] = sum(q * m for c, _, q, m, n, *_ in ITENS if c == cat)
        vals[f"G{r}"] = sum(q * n for c, _, q, m, n, *_ in ITENS if c == cat)
        r += 1
    vals[f"E{r}"] = sum(q * m for _, _, q, m, n, *_ in ITENS)
    vals[f"G{r}"] = sum(q * n for _, _, q, m, n, *_ in ITENS)
    tmp = pathlib.Path(tempfile.mkdtemp()) / "x.xlsx"
    with zipfile.ZipFile(caminho) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            dado = zi.read(it.filename)
            if it.filename == "xl/worksheets/sheet1.xml":
                xml = dado.decode()
                def troca(mo):
                    ref = mo.group(1)
                    return mo.group(0).replace("<v />", f"<v>{vals[ref]}</v>") if ref in vals else mo.group(0)
                xml, n = re.subn(r'<c r="([A-Z]+[0-9]+)"[^>]*><f>[^<]*</f><v />', troca, xml)
                dado = xml.encode()
            zo.writestr(it, dado)
    shutil.move(str(tmp), caminho)
    return vals


v = gravar_valores(OUT)
print("valores gravados:", len(v))
